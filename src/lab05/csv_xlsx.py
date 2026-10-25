import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

MIN_WIDTH = 8  # минимальная ширина колонки (в символах)
PADDING = 2  # небольшой запас, чтобы текст не прилипал к границе


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV в XLSX (библиотека openpyxl).

    Первая строка CSV - заголовок. Лист называется ``"Sheet1"``.
    Ширина каждой колонки = длина самого длинного текста в ней + запас,
    но не менее 8 символов. Значения записываются как текст (строки).

    Args:
        csv_path: путь к входному ``.csv`` файлу (UTF-8).
        xlsx_path: путь к выходному ``.xlsx`` файлу.

    Raises:
        FileNotFoundError: если CSV-файл не найден.
        ValueError: если расширения не ``.csv``/``.xlsx`` или CSV пустой
            (нет даже заголовка).
    """
    src, dst = Path(csv_path), Path(xlsx_path)
    if src.suffix.lower() != ".csv":
        raise ValueError(f"Неверный тип файла: ожидался .csv, получен {src.name!r}")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError(f"Неверный тип файла: ожидался .xlsx, получен {dst.name!r}")

    with open(src, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows or not any(cell.strip() for cell in rows[0]):
        raise ValueError("CSV-файл пуст или без заголовка")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)

    for col_idx in range(1, max(len(r) for r in rows) + 1):
        longest = max((len(r[col_idx - 1]) for r in rows if col_idx <= len(r)), default=0)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(MIN_WIDTH, longest + PADDING)

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)


if __name__ == "__main__":
    # Демонстрация (запуск из корня проекта: python -m src.lab05.csv_xlsx)
    csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")
    print("Готово: data/out/cities.xlsx")